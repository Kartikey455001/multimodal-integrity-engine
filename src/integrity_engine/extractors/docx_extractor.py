"""python-docx extraction. DOCX files do not provide reliable PDF-style page numbers."""

from __future__ import annotations

from pathlib import Path

from ..exceptions import OptionalDependencyError
from ..preprocessing.normalization import normalize_text
from ..schemas import Segment


def extract_docx(path: Path) -> tuple[list[Segment], list[str]]:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise OptionalDependencyError("DOCX extraction requires python-docx: pip install python-docx") from exc

    try:
        doc = DocxDocument(path)
    except Exception as exc:
        raise ValueError(f"Unreadable DOCX: {exc}") from exc

    segments: list[Segment] = []
    warnings = [
        "DOCX locations use 1-based paragraph or table-row identifiers; page numbers are not inferred without rendering."
    ]
    for i, paragraph in enumerate(doc.paragraphs, start=1):
        raw = paragraph.text or ""
        text = normalize_text(raw)
        if text:
            segments.append(
                Segment(
                    text=text,
                    location={"paragraph": i, "kind": "paragraph"},
                    original_text=raw,
                )
            )
    for table_index, table in enumerate(doc.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            raw = " | ".join(cell.text or "" for cell in row.cells)
            text = normalize_text(raw)
            if text:
                segments.append(
                    Segment(
                        text=text,
                        location={"table": table_index, "row": row_index, "kind": "table", "paragraph": None},
                        original_text=raw,
                    )
                )
    return segments, warnings
